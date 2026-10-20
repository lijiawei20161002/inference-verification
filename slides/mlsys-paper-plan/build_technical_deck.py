"""Build the BF16 technical supplement; shared layout copied from build_deck.py.

Run with a Python environment containing python-pptx and reportlab.
All plotted measurements are read from committed summaries. No GPU/API calls.
"""
from __future__ import annotations

import hashlib
import json
import re
import statistics
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
QA = HERE / "qa/technical"
PDF_DIR = ROOT / "output/pdf"
QA.mkdir(parents=True, exist_ok=True)
PDF_DIR.mkdir(parents=True, exist_ok=True)
PPTX = HERE / "inference-verification-technical-novelty.pptx"
PDF = PDF_DIR / "inference-verification-technical-novelty.pdf"
COMMIT = "57ceb5e"
CFP = "https://mlsys.org/Conferences/2027/CallForResearchPapers"
W, H = 960, 540

INK = "#172C40"
MUTED = "#586C7C"
BLUE = "#2864CF"
TEAL = "#087F82"
ORANGE = "#BA5B2B"
RED = "#B24143"
BG = "#F5F7FA"
PALE = "#E8EFF9"
TEAL_PALE = "#E4F2EF"
ORANGE_PALE = "#FAEEE4"
LINE = "#D9E1E9"
WHITE = "#FFFFFF"
DARK = "#142A3D"

for name, path in [
    ("CN", "/System/Library/Fonts/STHeiti Light.ttc"),
    ("CNBold", "/System/Library/Fonts/STHeiti Medium.ttc"),
    ("EN", "/System/Library/Fonts/Supplemental/Arial.ttf"),
    ("ENBold", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]:
    pdfmetrics.registerFont(TTFont(name, path))

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = "Inference Verification | BF16 与技术创新"
prs.core_properties.subject = "截至 2026-10-20 的归档结果、论文结构和待验证方案"
prs.core_properties.author = "inference-verification"
pdf = Canvas(str(PDF), pagesize=(W, H))
pdf.setTitle(prs.core_properties.title)
pdf.setAuthor("inference-verification")
pages, issues, evidence = [], [], {}
slide = None


def load(name, path):
    file = ROOT / path
    value = json.loads(file.read_text())
    evidence[name] = {"path": path, "sha256": hashlib.sha256(file.read_bytes()).hexdigest()}
    return value


q = load("quantization", "docs/results/real_quantization_l40s_20261018/pilot/summary.json")
v = load("serving", "docs/results/vllm_serving_l40s_20261019/summary.json")
c = load("cross_gpu", "docs/results/cross_gpu_replay_20261018/summary.json")
sched = load("scheduling", "docs/results/scheduling_l40s_20261019/summary.json")
assert q["detectors"]["token_difr"]["evaluation"]["nf4"]["rejected"] == 34
assert q["plan"]["calibration"] == 32 and q["plan"]["evaluation"] == 64
assert all(p["original_summary"]["measured_requests"] == 600 for p in v["phases"].values())
assert sched["verified_runs"] == 26 and sched["requests"] == 98
assert c["n_prompts"] == 80 and c["tokens_per_prompt"] == 256


def rgb(code):
    return RGBColor.from_string(code.lstrip("#"))


def rect(x, y, w, h, fill, stroke=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Pt(x), Pt(y), Pt(w), Pt(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if stroke:
        shape.line.color.rgb = rgb(stroke)
        shape.line.width = Pt(0.7)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    pdf.setFillColor(HexColor(fill))
    pdf.setStrokeColor(HexColor(stroke or fill))
    pdf.setLineWidth(0.7)
    if radius:
        pdf.roundRect(x, H-y-h, w, h, 8, fill=1, stroke=int(bool(stroke)))
    else:
        pdf.rect(x, H-y-h, w, h, fill=1, stroke=int(bool(stroke)))
    return shape


def line(x1, y1, x2, y2, color=LINE, width=1):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    pdf.setStrokeColor(HexColor(color))
    pdf.setLineWidth(width)
    pdf.line(x1, H-y1, x2, H-y2)


def dot(x, y, r, color):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Pt(x-r), Pt(y-r), Pt(2*r), Pt(2*r))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    pdf.setFillColor(HexColor(color))
    pdf.circle(x, H-y, r, stroke=0, fill=1)


def wrap(value, width, font, size):
    result = []
    for paragraph in str(value).split("\n"):
        if not paragraph:
            result.append("")
            continue
        tokens = re.findall(r"[A-Za-z0-9_./%+′≤≥=-]+|[^\x00-\x7F]|\s+|.", paragraph)
        current = ""
        for token in tokens:
            if current and pdfmetrics.stringWidth(current+token, font, size) > width:
                result.append(current.rstrip())
                current = token.lstrip()
            else:
                current += token
        result.append(current.rstrip())
    return result


def text(x, y, w, h, value, size=18, color=INK, bold=False, align="left",
         english=False, leading=1.22, link_url=None):
    font = ("EN" if english else "CN") + ("Bold" if bold else "")
    lines = wrap(value, w-2, font, size)
    if len(lines)*size*leading > h+1:
        issues.append(f"p{len(pages)} overflow: {value[:60]} ({len(lines)} lines in {h} pt)")
    for value_line in lines:
        if pdfmetrics.stringWidth(value_line, font, size) > w+0.5:
            issues.append(f"p{len(pages)} width: {value_line}")
    if x < 0 or y < 0 or x+w > W+1 or y+h > H+1:
        issues.append(f"p{len(pages)} out of page: {value[:60]}")
    sh = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, value_line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size*leading)
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        run = p.add_run()
        run.text = value_line
        run.font.name = "Arial" if english else "Heiti SC"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        ea = OxmlElement("a:ea")
        ea.set("typeface", "Heiti SC")
        run._r.get_or_add_rPr().append(ea)
        if link_url:
            run.hyperlink.address = link_url
        pdf.setFont(font, size)
        pdf.setFillColor(HexColor(color))
        baseline = H-y-size*0.88-i*size*leading
        if align == "center":
            pdf.drawCentredString(x+w/2, baseline, value_line)
        elif align == "right":
            pdf.drawRightString(x+w, baseline, value_line)
        else:
            pdf.drawString(x, baseline, value_line)
    if link_url:
        pdf.linkURL(link_url, (x, H-y-h, x+w, H-y), thickness=0)
    return sh


def new_page(section, title, subtitle, sources, notes, status="已测量", dark=False):
    global slide
    if pages:
        pdf.showPage()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pages.append({"page": len(pages)+1, "section": section, "title": title,
                  "status": status, "sources": sources, "notes": notes})
    rect(0, 0, W, H, DARK if dark else BG)
    ink, muted = (WHITE, "#B7C8D6") if dark else (INK, MUTED)
    text(36, 22, 700, 18, f"INFERENCE VERIFICATION  /  {section}", 11, "#8BCDC8" if dark else TEAL,
         True, english=True)
    text(750, 22, 172, 18, status, 11, muted, align="right")
    text(36, 60, 888, 77, title, 29, ink, True)
    if subtitle:
        text(36, 139, 888, 45, subtitle, 16, muted)
    line(36, 494, 924, 494, "#3A5063" if dark else LINE)
    source_text = "  |  ".join(sources)
    text(36, 502, 850, 15, source_text, 9, muted, english=True)
    text(36, 520, 830, 13, f"2026-10-20  ·  archived through 2026-10-19  ·  {COMMIT}", 8, muted, english=True)
    text(890, 512, 34, 18, f"{len(pages):02d}", 12, muted, align="right", english=True)
    slide.notes_slide.notes_text_frame.text = (
        title+"\n\n"+notes+"\n\nSources:\n"+"\n".join(sources)+
        "\nLocal source manifest: slides/mlsys-paper-plan/qa/technical/evidence.json"
    )
    return slide


def banner(value, y=447, fill=PALE, color=BLUE, size=17, h=33):
    rect(36, y, 888, h, fill, radius=True)
    text(48, y+6, 864, h-8, value, size, color, True)


def note(x, y, w, value, size=12):
    text(x, y, w, 43, value, size, MUTED)


def contrast(phase, config, length, concurrency):
    return next(z for z in v["phases"][phase]["contrasts"]
                if z["config"] == config and z["length"] == length and z["concurrency"] == concurrency)


def errorbar(x, y, lo, hi, width, maxval, color):
    a, b = x+lo/maxval*width, x+hi/maxval*width
    line(a, y, b, y, color, 1.4)
    line(a, y-3, a, y+3, color, 1.4)
    line(b, y-3, b, y+3, color, 1.4)


SOURCES = {
    "DiFR": "https://arxiv.org/abs/2511.20621",
    "RUT": "https://arxiv.org/abs/2506.06975",
    "TOPLOC": "https://arxiv.org/abs/2501.16007",
    "Batch invariance": "https://thinkingmachines.ai/blog/defeating-nondeterminism-in-llm-inference/",
    "LLM-42": "https://arxiv.org/abs/2601.17768",
    "Hawkeye": "https://proceedings.mlsys.org/paper_files/paper/2026/hash/e217c271a57c365a246b0ad39e668ba8-Abstract-Conference.html",
    "Bit-Exact": "https://arxiv.org/abs/2606.00279",
    "Taming Bitwise": "https://arxiv.org/abs/2609.11356",
    "Tensor Core Models": "https://arxiv.org/abs/2512.07004",
    "INT8 Conformance": "https://arxiv.org/abs/2609.00363",
    "FlashInfer": "https://docs.flashinfer.ai/api/attention.html",
    "vLLM": "https://docs.vllm.ai/en/latest/features/batch_invariance/",
    "Sarathi-Serve": "https://www.usenix.org/conference/osdi24/presentation/agrawal",
    "FlashAttention-4": "https://arxiv.org/abs/2603.05451",
}
tail = load("score_tails", "slides/mlsys-paper-plan/qa/technical/score-tail-diagnostic.json")


def citations(keys, y=467):
    widths = [min(200, max(80, len(k)*7+10)) for k in keys]
    x = 36
    for key, width in zip(keys, widths):
        text(x, y, width, 18, key, 10.5, BLUE, english=True, link_url=SOURCES[key])
        x += width + 14


def source_notes(keys, value):
    return value + "\n\n" + "\n".join(f"{k}: {SOURCES[k]}" for k in keys)


# 1. Correct the claim before proposing a new method.
new_page(
    "TECHNICAL POSITION", "BF16 还未解决。\n核心创新需要重新收敛。",
    "已有实证基础；“校准 + 多硬件 + 接 vLLM”不足以成为新的核心贡献。",
    ["local archives + primary literature through 2026-10-20"],
    source_notes(["DiFR", "LLM-42", "Bit-Exact"],
        "这是对上一版创新性判断的修正。现有工作没有完成新验证算法；本 deck 提出可证伪方向。"
        "必须区分重复性、跨执行合法性、检测能力和证明物理精度。"),
    status="研究判断 / 尚无新算法结果",
)
for i, (n, title, body) in enumerate([
    ("01", "运行可重复", "同一路径复跑一致\n不等于不同 shape 一致"),
    ("02", "诚实 BF16 通过", "跨 kernel / GPU / 调度\n误报率仍需独立验证"),
    ("03", "低精度可检出", "真实 INT8 仅检出 5/64\n不够支撑广泛可靠性"),
]):
    x = 36+i*300
    rect(x, 205, 288, 184, WHITE, radius=True)
    text(x+16, 222, 250, 34, n, 24, TEAL, True, english=True)
    text(x+16, 269, 250, 30, title, 21, INK, True)
    text(x+16, 315, 250, 62, body, 17, MUTED)
banner("合理目标：检验明确契约下的行为一致性；不能由 token 无条件证明物理精度。", 408, size=16)
citations(["DiFR", "LLM-42", "Bit-Exact"])

# 2. New reproducible evidence.
new_page(
    "NEW OFFLINE DIAGNOSTIC", "49 个封顶位置，贡献了\n诚实 BF16 总分的 98.4%。",
    "64 条 × 128 tokens 的原评估切分。少量极端分数值得优先归因。",
    ["real_quantization_l40s_20261018/pilot/scores.npz", "qa/technical/score-tail-diagnostic.json"],
    "每组8192个token。封顶值30也可能来自有限大margin，尚不能从分数证明候选集不一致。"
    "所有这些位置同时触发本地rank cap50。图为总分质量占比，不是FPR或TPR；"
    "BF16含封顶序列32/64也不是误报32/64。不能直接删除封顶事件宣称新检测性能。",
    status="新离线分析 / 归因待验证",
)
for i, (arm, label, col) in enumerate([("bf16", "Honest BF16", TEAL), ("nf4", "NF4", BLUE), ("int8", "INT8", ORANGE)]):
    a = tail["arms"][arm]
    y = 204+i*65
    text(36, y+4, 150, 27, label, 18, col, True, english=True)
    rect(195, y+7, 500, 22, LINE)
    rect(195, y+7, 500*a["capped_score_mass_fraction"], 22, col)
    text(710, y+1, 205, 30, f'{a["capped_score_mass_fraction"]:.2%}', 23, col, True, english=True)
    text(195, y+34, 650, 20, f'封顶位置 {a["capped_scores"]:,} / 8,192；其余位置贡献剩余总分。', 12, MUTED)
text(36, 407, 888, 31, "代码中：候选被 top-k/top-p 排除 → 分数直接记为 30。", 19, INK, True)
banner("线索：过滤边界可能放大小误差。下一步要记录原始 logits 与两端候选集。", 447, size=16)

# 3. Causal chain, not generic numerical noise.
new_page(
    "MECHANISM", "先把协议差异与硬件误差分开。",
    "BF16 dtype 不是完整数值契约；输出差异可能在多个环节产生。",
    ["ivgym/sampling.py", "ivgym/verifiers.py", "experiments/exp_real_quantization_gpu.py"],
    "本地provider float32/verifier float64过滤。NumPy RNG不复现vLLM随机流。"
    "top-p注释包括crossing token，代码cum<=p；[.6,.3,.1],p=.7当前只保留第一项。"
    "这些是需要对齐的控制变量，不足以认定是历史pilot封顶的主因。"
    "合成top-k例子来自分析脚本，非真实模型事件。本次不改采样器。",
    status="代码核对 + 机制假设",
)
for i, (title, desc, col) in enumerate([
    ("执行路径", "GEMM / split-KV\n归约、cast、SFU", TEAL),
    ("数值误差", "冻结 prefix 后\n定位 logit 变化", BLUE),
    ("过滤边界", "top-k / top-p\n候选集离散跳变", ORANGE),
    ("采样分数", "共享随机数下\n从 0 跳到封顶", RED),
]):
    x = 36+i*224
    rect(x, 199, 212, 111, WHITE, radius=True)
    text(x+12, 211, 185, 28, title, 21, col, True)
    text(x+12, 252, 188, 48, desc, 15, MUTED)
    if i < 3:
        text(x+213, 240, 12, 26, "›", 19, MUTED)
rect(36, 332, 888, 96, PALE, radius=True)
text(51, 345, 854, 26, "先做同 logits 的 golden replay，再讨论 GPU 机制。", 19, BLUE, True)
text(51, 382, 854, 38, "逐项对齐 float32/64 过滤、top-p 边界约定、tie-breaking 和 RNG。\n当前 top-p 注释与代码不一致；修订协议后必须重新校准。", 14, MUTED)
banner("本次未修改采样器。若只是协议问题，修复基线；不包装成硬件创新。", 447, size=16)

# 4. Verification baselines and access conditions.
new_page(
    "RELATED WORK / VERIFICATION", "对比必须包括最接近的方法，\n并匹配可获得的信息。",
    "本地 token_toploc 是 token rank 基线，不能当作原版 TOPLOC 或 RUT。",
    ["DiFR (2025)", "TOPLOC (2025)", "RUT (2025 / ICLR 2026)"],
    source_notes(["DiFR", "TOPLOC", "RUT"],
        "DiFR已包含合法配置混合校准、likelihood与Gaussian Monte Carlo变体。"
        "RUT需要本地参考样本，应计入成本。TOPLOC/Activation-DiFR需要provider提供激活信息。"
        "不同信息条件单列结果，不混淆为一个公平排行榜。"),
    status="文献对比 / 必做基线",
)
cols = [(36, 178), (224, 310), (546, 378)]
for (x, width), val in zip(cols, ["方法 / 信息", "已经做到", "我们必须新增什么"]):
    rect(x, 193, width, 30, DARK)
    text(x+9, 201, width-18, 18, val, 13, WHITE, True)
rows = [
    ("DiFR\n共享 RNG", "margin / likelihood / 高斯模拟\n合法配置混合校准 + vLLM", "比通用噪声更紧的采样边界建模\n同误报下更好的能力—成本"),
    ("RUT\n无共享 RNG", "以本地参考样本校准秩统计\n面向黑盒行为分布", "单列 seed-free 对照\n参考采样不是免费"),
    ("TOPLOC / Act-DiFR\n需要 provider 激活", "中间状态的压缩指纹比较", "单列合作式验证条件\n不能用额外状态暗中增强 token-only"),
]
for i, values in enumerate(rows):
    y=231+i*65
    for (x, width), value in zip(cols, values):
        rect(x, y, width, 59, WHITE)
        text(x+9, y+10, width-18, 47, value, 13.2, INK)
banner("“混合校准 + 新打分 + 接 serving”不足；要证明机制建模带来独立收益。", 438, size=16, h=29)
citations(["DiFR", "RUT", "TOPLOC"], 474)

# 5. Hardware frontier.
new_page(
    "RELATED WORK / HARDWARE", "2026 年工作已深入到\n归约树、指令语义与精确回放。",
    "硬件知识应成为方法的约束和干预手段；仅指出浮点非结合性已经不够。",
    ["primary papers; June / September works below are preprints"],
    source_notes(["LLM-42", "Hawkeye", "Bit-Exact", "Taming Bitwise", "Tensor Core Models"],
        "这些工作提供不同层级的数值控制。LLM-42的确定性服务与外部恶意provider审计不同。"
        "Hawkeye针对Tensor Core矩阵运算；Bit-Exact在测试配置上精确回放模型块，不代表所有模型和未知引擎。"
        "Taming已涉及kernel等价类，不能把等价类概念当作新贡献。"),
    status="最新技术定位",
)
hardware_rows = [
    ("LLM-42 · 2026", "快速解码 + 固定形状窗口验证 / 回滚", "固定窗口回放本身已有先例", "LLM-42"),
    ("Hawkeye · MLSys 2026", "Tensor Core matmul 数值建模 / CPU 回放", "从实际归约与舍入语义入手", "Hawkeye"),
    ("Bit-Exact · 2026-06", "Tensor Core、SFU、cast 与模型块精确回放", "对照完整回放的执行信息与成本", "Bit-Exact"),
    ("Taming Bitwise · 2026-09", "归约描述、编译约束、静态验证 / 等价类", "机会在跨层到采样的组合", "Taming Bitwise"),
]
for i, (label, desc, implication, key) in enumerate(hardware_rows):
    y=191+i*59
    rect(36, y, 888, 53, WHITE, radius=True)
    text(48, y+10, 267, 33, label, 15, BLUE, True, link_url=SOURCES[key])
    text(325, y+8, 347, 37, desc, 13.8, INK)
    text(689, y+8, 223, 37, implication, 13, MUTED)
banner("候选空白：把底层数值约束组合成采样决定，并让按需验证确实省成本。", 438, size=16, h=29)
citations(["Tensor Core Models", "Batch invariance"], 474)

# 6. Concrete algorithm, explicit conditional guarantee.
new_page(
    "PROPOSED ALGORITHM", "从一个参考分数，走向\n“合法数值变化能否产生这个 token？”",
    "候选方案：考虑候选集变化的数值一致性验证。尚未实现或验证收益。",
    ["proposal", "planning/BF16_TECHNICAL_NOVELTY.md"],
    "固定prefix、共享随机数与一个合法类别；无过滤或已证明支持集固定时，z_j=l_j+T*g_j。"
    "有效区间[L_j,U_j]下，若始终可用的竞争token满足L_j>U_y则claimed y不可能胜出。"
    "top-k/p支持集变化必须联合考虑；区间重叠并不证明可达。"
    "经验包络只能提供经验校准，不等于形式化certificate。未知执行类别须全部覆盖。",
    status="算法设计 / 待验证",
)
rect(36, 193, 422, 226, WHITE, radius=True)
text(52, 208, 390, 27, "固定候选集：先解决可控子问题", 19, TEAL, True)
text(52, 246, 390, 28, "z[j] = logit[j] + T * g[j]", 19, INK, english=True)
text(52, 280, 390, 24, "合法变化下，z[j] 位于区间 [L[j], U[j]]。", 15, MUTED)
text(52, 319, 390, 34, "若某个合法竞争者： L[j] > U[y]", 19, BLUE, True)
text(52, 364, 390, 47, "则 claimed token y 不可能获胜。\n区间重叠 → 仍有歧义，不能直接判通过。", 16, MUTED)
rect(474, 193, 450, 226, WHITE, radius=True)
text(490, 208, 418, 27, "难点：top-k / top-p 的离散边界", 19, ORANGE, True)
for j, value in enumerate([
    "同时约束 logit 与候选成员资格",
    "保留跨层误差相关性，避免区间膨胀",
    "剪枝不可能的候选集；歧义触发复算",
    "无法收紧时输出 inconclusive",
]):
    dot(496, 261+j*40, 3, ORANGE)
    text(509, 249+j*40, 394, 34, value, 16, INK)
banner("核心区别必须来自边界建模与误差约束；DiFR 已有 Gaussian Monte Carlo。", 447, size=15.8)

# 7. Engineering architecture and cost.
new_page(
    "PROPOSED SYSTEM", "用硬件机制缩小不确定性，\n把复算花在能消除歧义的位置。",
    "这不是“多加几次 replay”：必须测出相同错误率下的验证成本收益。",
    ["proposal", "Hawkeye / LLM-42 / Bit-Exact / Taming Bitwise"],
    source_notes(["Hawkeye", "LLM-42", "Bit-Exact", "Taming Bitwise"],
        "合法执行目录必须独立定义；provider自报元信息不是可信证据。"
        "同类别p值有效且真实类别被覆盖时max_c p_c对有限并集保守，构造本身不是创新。"
        "动态执行须按合法轨迹建模，不能逐token任意选择最有利类别。"
        "选择性回放可能仍需整段prefix/KV重建，所有成本必须计入。"),
    status="系统设计 / 成本待测",
)
steps=[
    ("01", "合法执行目录", "GPU / kernel / shape\n归约、cast、采样协议"),
    ("02", "一次参考 pass", "估计数值不确定性\n判定支持集是否稳定"),
    ("03", "按需精确回放", "歧义位置 / 窗口\n复用可信 prefix 状态"),
    ("04", "冻结规则评估", "异常 / 未排除 / 拒判\n独立校准与测试"),
]
for i, (n, title, body) in enumerate(steps):
    x=36+i*224
    rect(x, 199, 212, 153, WHITE, radius=True)
    text(x+13, 212, 184, 24, n, 17, TEAL, True, english=True)
    text(x+13, 249, 184, 26, title, 19, INK, True)
    text(x+13, 292, 184, 49, body, 14.5, MUTED)
text(36, 373, 888, 59, "成本陷阱：回放一个 token 也可能要重建整个 prefix / KV。\n信任陷阱：未知执行类别不能只凭 provider 自报来缩小。", 18, ORANGE, True)
banner("要同时报告检测能力、拒判率、参考 GPU 时间，以及缓存与校准成本。", 447, size=16)

# 8. Serving and timing stays grounded in mechanisms.
new_page(
    "SYSTEMS CONTROLS", "把系统知识变成干预实验；\nclock 暂作辅助证据。",
    "现有调度与缓存结果有价值，但尚未证明跨部署可辨识的 timing 签名。",
    ["vLLM / FlashInfer docs", "Sarathi-Serve (2024)", "FlashAttention-4 (2026)"],
    source_notes(["vLLM", "FlashInfer", "Sarathi-Serve", "FlashAttention-4"],
        "vLLM batch invariance的文档范围不等于跨硬件保证。FlashInfer固定split需要对应backend。"
        "FA4面向Blackwell，不能将吞吐结论搬到L40S。"
        "现有调度干预同时可能改变编译shape；没有DRAM/L2/occupancy计数器。API并发不等于实际batch。"
        "量化本地pilot更慢，不能假设半字节数必然更快；恶意延迟填充限制clock可辨识性。"),
    status="机制对照 / 待补因果证据",
)
system_rows = [
    ("vLLM / FlashInfer", "固定 batch 归约或 split-KV 合并", "数值差异能否随干预消失？", "FlashInfer"),
    ("Sarathi-Serve", "chunked prefill 与 decode 交错", "13 → 4 次混合迭代是否解释早窗？", "Sarathi-Serve"),
    ("FlashAttention-4", "TMEM、异步 MMA、softmax / exp 瓶颈", "避免只用逻辑字节数解释 timing", "FlashAttention-4"),
]
for i,(label,mechanism,question,key) in enumerate(system_rows):
    y=197+i*72
    rect(36,y,888,64,WHITE,radius=True)
    text(49,y+11,244,44,label,17,BLUE,True,link_url=SOURCES[key])
    text(304,y+11,323,44,mechanism,15,INK)
    text(642,y+11,267,44,question,14,MUTED)
banner("尚缺硬件计数器与充分的单因素控制；量化加速、纯带宽模型都不能预设。", 438, size=15.6, h=29)
citations(["vLLM", "FlashInfer", "FlashAttention-4"], 474)

# 9. Decisive experiments instead of another broad sweep.
new_page(
    "GO / NO-GO EXPERIMENTS", "先做三组能决定方向的实验。",
    "收益必须在相同信息条件、相同误报约束和独立测试上成立。",
    ["proposed experiments; no new GPU jobs run"],
    "先协议归因，再机制干预，最后冻结方法并进行独立评价。"
    "FPR按合法条件和独立序列/session报告，多个条件应处理同时覆盖。"
    "不能用token数代替独立样本数。拒判率必须报告，避免把不确定样本全部移除得到虚假收益。"
    "检测对照包含DiFR margin/likelihood/MC、混合校准和多配置完整回放；RUT与激活方法单列信息条件。",
    status="下一轮实验计划",
)
experiments=[
    ("P0", "定位 49 个 BF16 封顶事件", "记录两端 logits、候选集、边界、margin；先对齐采样协议。",
     "若主要是协议差异 → 修复并重跑基线，不宣称硬件创新。", TEAL),
    ("P1", "冻结输入，只干预一个数值机制", "GEMM shape / 归约 / attention split；用确定性路径作对照。",
     "若误差与具体机制关联很弱 → 不扩展昂贵的全栈建模。", BLUE),
    ("P2", "测能力—成本曲线与消融", "DiFR 各变体、完整回放、候选方法；真实 NF4 / INT8。",
     "若同时抹掉攻击信号，或复算不省成本 → 该主线失败。", ORANGE),
]
for i,(p,title,design,gate,col) in enumerate(experiments):
    y=193+i*82
    rect(36,y,888,75,WHITE,radius=True)
    text(48,y+24,49,25,p,18,col,True,english=True)
    text(109,y+9,324,52,title,19,INK,True)
    text(442,y+9,465,31,design,13,MUTED)
    text(442,y+43,465,27,gate,12.8,col)
banner("指标：逐条件 FPR 上界、TPR、拒判率、审计 GPU 时间；新语料与独立 session。", 454, size=15.4, h=29)

# 10. Paper structure reflecting a concrete, falsifiable thesis.
new_page(
    "REVISED PAPER THESIS", "论文围绕一个问题写：\n能否既容忍合法 BF16，又保留量化信号？",
    "暂定方向：Sampling-Aware Verification of BF16 Inference under Execution Variability",
    ["planning/BF16_TECHNICAL_NOVELTY.md", "technical supplement, 2026-10-20"],
    "标题为建议，尚无完成的方法和正结果。"
    "写作结构：Problem与契约1页，机制1.5页，方法2.5页，系统1页，评估3页，局限/相关工作1页，"
    "仅为十页正文的分配建议。若关键命题失败，应收窄为经过充分验证的机制/边界研究，不补写成功结论。",
    status="论文主线 / 待证命题",
)
for i,(label,title,body,col) in enumerate([
    ("OBSERVATION","解释误报的来源","协议控制 → 过滤边界 → 硬件机制\n不要只展示“不同 GPU 有差异”",TEAL),
    ("METHOD","提出新的判据与实现","候选集感知 + 受约束误差模型\n歧义驱动复算，明确成本与适用范围",BLUE),
    ("EVIDENCE","验证独立增益","固定 FPR，比 TPR / GPU 时间\n真实量化 + 合法执行变化 + 消融",ORANGE),
]):
    x=36+i*300
    rect(x,200,288,196,WHITE,radius=True)
    text(x+15,216,257,24,label,13,col,True,english=True)
    text(x+15,257,257,31,title,22,INK,True)
    text(x+15,307,257,75,body,16,MUTED)
banner("目前可以写问题和机制；核心方法与收益必须靠下一轮实验决定，不能预写结论。", 416, size=15.5)
text(36,465,888,18,"已交付：技术评审、最新文献对照、离线分数诊断与 10 页可编辑 slides。",12,MUTED)

pdf.save()
prs.save(PPTX)
evidence["literature"] = {"checked": "2026-10-20", "sources": SOURCES}
for name, value in [
    ("evidence.json", evidence),
    ("slide-index.json", pages),
    ("layout-validation.json", {
        "slides": len(pages), "overflow_issues": issues,
        "renderer": "editable PPTX + ReportLab PDF, shared geometry",
        "pptx_rendering_limit": "PowerPoint/Keynote native rendering not checked",
    }),
]:
    (QA/name).write_text(json.dumps(value, indent=2, ensure_ascii=False)+"\n")
if issues:
    raise SystemExit("\n".join(issues))
print(f"Built {len(pages)} slides with no measured text overflow.\n{PPTX}\n{PDF}")
