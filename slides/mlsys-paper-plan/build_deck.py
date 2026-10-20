"""Build an editable, evidence-backed Chinese MLSys paper-plan deck and matching PDF.

Run with a Python environment containing python-pptx and reportlab.
All plotted measurements are read from committed summaries. No GPU/API calls.
"""
from __future__ import annotations

import hashlib
import json
import re
import statistics
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
QA = HERE / "qa"
PDF_DIR = ROOT / "output/pdf"
QA.mkdir(parents=True, exist_ok=True)
PDF_DIR.mkdir(parents=True, exist_ok=True)
PPTX = HERE / "inference-verification-mlsys.pptx"
PDF = PDF_DIR / "inference-verification-mlsys.pdf"
COMMIT = "57ceb5e"
CFP = "https://mlsys.org/Conferences/2027/CallForResearchPapers"
W, H = 960, 540

INK = "#172C40"
MUTED = "#586C7C"
BLUE = "#2864CF"
TEAL = "#087F82"
ORANGE = "#BA5B2B"
RED = "#B24143"
BG = "#F5F7FA"
PALE = "#E8EFF9"
TEAL_PALE = "#E4F2EF"
ORANGE_PALE = "#FAEEE4"
LINE = "#D9E1E9"
WHITE = "#FFFFFF"
DARK = "#142A3D"

for name, path in [
    ("CN", "/System/Library/Fonts/STHeiti Light.ttc"),
    ("CNBold", "/System/Library/Fonts/STHeiti Medium.ttc"),
    ("EN", "/System/Library/Fonts/Supplemental/Arial.ttf"),
    ("ENBold", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]:
    pdfmetrics.registerFont(TTFont(name, path))

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = "Inference Verification | 实验进展与 MLSys 论文主线"
prs.core_properties.subject = "截至 2026-10-20 的归档结果、论文结构和待验证方案"
prs.core_properties.author = "inference-verification"
pdf = Canvas(str(PDF), pagesize=(W, H))
pdf.setTitle(prs.core_properties.title)
pdf.setAuthor("inference-verification")
pages, issues, evidence = [], [], {}
slide = None


def load(name, path):
    file = ROOT / path
    value = json.loads(file.read_text())
    evidence[name] = {"path": path, "sha256": hashlib.sha256(file.read_bytes()).hexdigest()}
    return value


q = load("quantization", "docs/results/real_quantization_l40s_20261018/pilot/summary.json")
v = load("serving", "docs/results/vllm_serving_l40s_20261019/summary.json")
c = load("cross_gpu", "docs/results/cross_gpu_replay_20261018/summary.json")
sched = load("scheduling", "docs/results/scheduling_l40s_20261019/summary.json")
assert q["detectors"]["token_difr"]["evaluation"]["nf4"]["rejected"] == 34
assert q["plan"]["calibration"] == 32 and q["plan"]["evaluation"] == 64
assert all(p["original_summary"]["measured_requests"] == 600 for p in v["phases"].values())
assert sched["verified_runs"] == 26 and sched["requests"] == 98
assert c["n_prompts"] == 80 and c["tokens_per_prompt"] == 256


def rgb(code):
    return RGBColor.from_string(code.lstrip("#"))


def rect(x, y, w, h, fill, stroke=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Pt(x), Pt(y), Pt(w), Pt(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if stroke:
        shape.line.color.rgb = rgb(stroke)
        shape.line.width = Pt(0.7)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    pdf.setFillColor(HexColor(fill))
    pdf.setStrokeColor(HexColor(stroke or fill))
    pdf.setLineWidth(0.7)
    if radius:
        pdf.roundRect(x, H-y-h, w, h, 8, fill=1, stroke=int(bool(stroke)))
    else:
        pdf.rect(x, H-y-h, w, h, fill=1, stroke=int(bool(stroke)))
    return shape


def line(x1, y1, x2, y2, color=LINE, width=1):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    pdf.setStrokeColor(HexColor(color))
    pdf.setLineWidth(width)
    pdf.line(x1, H-y1, x2, H-y2)


def dot(x, y, r, color):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Pt(x-r), Pt(y-r), Pt(2*r), Pt(2*r))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    pdf.setFillColor(HexColor(color))
    pdf.circle(x, H-y, r, stroke=0, fill=1)


def wrap(value, width, font, size):
    result = []
    for paragraph in str(value).split("\n"):
        if not paragraph:
            result.append("")
            continue
        tokens = re.findall(r"[A-Za-z0-9_./%+′≤≥=-]+|[^\x00-\x7F]|\s+|.", paragraph)
        current = ""
        for token in tokens:
            if current and pdfmetrics.stringWidth(current+token, font, size) > width:
                result.append(current.rstrip())
                current = token.lstrip()
            else:
                current += token
        result.append(current.rstrip())
    return result


def text(x, y, w, h, value, size=18, color=INK, bold=False, align="left",
         english=False, leading=1.22, link_url=None):
    font = ("EN" if english else "CN") + ("Bold" if bold else "")
    lines = wrap(value, w-2, font, size)
    if len(lines)*size*leading > h+1:
        issues.append(f"p{len(pages)} overflow: {value[:60]} ({len(lines)} lines in {h} pt)")
    for value_line in lines:
        if pdfmetrics.stringWidth(value_line, font, size) > w+0.5:
            issues.append(f"p{len(pages)} width: {value_line}")
    if x < 0 or y < 0 or x+w > W+1 or y+h > H+1:
        issues.append(f"p{len(pages)} out of page: {value[:60]}")
    sh = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, value_line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size*leading)
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        run = p.add_run()
        run.text = value_line
        run.font.name = "Arial" if english else "Heiti SC"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        ea = OxmlElement("a:ea")
        ea.set("typeface", "Heiti SC")
        run._r.get_or_add_rPr().append(ea)
        if link_url:
            run.hyperlink.address = link_url
        pdf.setFont(font, size)
        pdf.setFillColor(HexColor(color))
        baseline = H-y-size*0.88-i*size*leading
        if align == "center":
            pdf.drawCentredString(x+w/2, baseline, value_line)
        elif align == "right":
            pdf.drawRightString(x+w, baseline, value_line)
        else:
            pdf.drawString(x, baseline, value_line)
    if link_url:
        pdf.linkURL(link_url, (x, H-y-h, x+w, H-y), thickness=0)
    return sh


def new_page(section, title, subtitle, sources, notes, status="已测量", dark=False):
    global slide
    if pages:
        pdf.showPage()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pages.append({"page": len(pages)+1, "section": section, "title": title,
                  "status": status, "sources": sources, "notes": notes})
    rect(0, 0, W, H, DARK if dark else BG)
    ink, muted = (WHITE, "#B7C8D6") if dark else (INK, MUTED)
    text(36, 22, 700, 18, f"INFERENCE VERIFICATION  /  {section}", 11, "#8BCDC8" if dark else TEAL,
         True, english=True)
    text(750, 22, 172, 18, status, 11, muted, align="right")
    text(36, 60, 888, 77, title, 29, ink, True)
    if subtitle:
        text(36, 139, 888, 45, subtitle, 16, muted)
    line(36, 494, 924, 494, "#3A5063" if dark else LINE)
    source_text = "  |  ".join(sources)
    text(36, 502, 850, 15, source_text, 9, muted, english=True)
    text(36, 520, 830, 13, f"2026-10-20  ·  archived through 2026-10-19  ·  {COMMIT}", 8, muted, english=True)
    text(890, 512, 34, 18, f"{len(pages):02d}", 12, muted, align="right", english=True)
    slide.notes_slide.notes_text_frame.text = (
        title+"\n\n"+notes+"\n\nSources:\n"+"\n".join(sources)+
        "\nLocal source manifest: slides/mlsys-paper-plan/qa/evidence.json"
    )
    return slide


def banner(value, y=447, fill=PALE, color=BLUE, size=17, h=33):
    rect(36, y, 888, h, fill, radius=True)
    text(48, y+6, 864, h-8, value, size, color, True)


def note(x, y, w, value, size=12):
    text(x, y, w, 43, value, size, MUTED)


def contrast(phase, config, length, concurrency):
    return next(z for z in v["phases"][phase]["contrasts"]
                if z["config"] == config and z["length"] == length and z["concurrency"] == concurrency)


def errorbar(x, y, lo, hi, width, maxval, color):
    a, b = x+lo/maxval*width, x+hi/maxval*width
    line(a, y, b, y, color, 1.4)
    line(a, y-3, a, y+3, color, 1.4)
    line(b, y-3, b, y+3, color, 1.4)


# 1. Outcome first: credible progress, bounded maturity.
new_page(
    "RESEARCH UPDATE / MLSYS",
    "实验已走到 serving 机制层。\n论文应回答：推理验证何时可信？",
    "建议题目  IVGym: Inference Verification under Serving Variability",
    ["real_quantization_l40s_20261018", "vllm_serving_l40s_20261019", "scheduling_l40s_20261019"],
    "开场：截至10月20日本地提交57ceb5e，最新归档测量为10月19日。已有真实权重量化、"
    "生产 serving 引擎及调度追踪证据，足以开始围绕系统问题写作。这里的成熟度评价是写作建议，"
    "不是投稿录用判断。1,200是两轮同机同输入loopback请求，26是另一个诊断campaign，不能合并为独立部署样本。"
    "尚无经独立校准和测试验证的联合token/clock检测器。",
    status="进展 + 写作建议", dark=True,
)
for x, big, label, detail in [
    (36, "34 / 64", "真实 NF4 被检出", "token-DiFR；128 tokens / 序列"),
    (338, "1,200", "vLLM 正式测量请求", "600 pilot + 600 同输入时间复测"),
    (640, "26 / 26", "调度诊断运行通过验证", "Nsight / PyTorch / 无 profiler"),
]:
    rect(x, 216, 284, 132, "#203D52", radius=True)
    text(x+17, 230, 250, 47, big, 36, "#8DDFD1", True, english=True)
    text(x+17, 284, 250, 25, label, 18, WHITE, True)
    text(x+17, 316, 250, 23, detail, 12, "#B7C8D6")
text(40, 382, 880, 75,
     "核心判断：机制证据已形成；部署误报率、跨条件泛化和验证成本尚未闭环。\n"
     "主线建议：用真实 serving 变化定义检测边界，再验证有适用条件的审计方法。",
     20, WHITE, leading=1.4)


# 2. Concrete archive status; do not count all campaigns as one joint study.
new_page(
    "01 / EXPERIMENT STATUS",
    "最新进展：六条证据链已归档，联合验证仍待完成",
    "“完成”指该实验协议完成；各项测量的模型、样本与统计单位不同。",
    ["docs/README.md", "experiments/README.md", "docs/results/*_20261016..19.md"],
    "H100八项基线和A100四项基线独立运行；跨GPU配对比较直到10月18日才真正完成。"
    "真实量化的0/64 BF16误报不是部署保证。vLLM两轮复用了同样的prompt和配置顺序。"
    "调度主格只有两个重叠文本offset，辅助格一个。以上不能拼成同请求的token+clock评估。"
    "8月paper.tex与README部分叙述落后于归档，本deck优先使用最新实验报告。",
)
cols = [36, 221, 554, 924]
for a, b in zip(cols[:-1], cols[1:]):
    rect(a, 196, b-a, 29, INK)
for x, w, val in [(46, 170, "实验 / 日期"), (231, 310, "已完成的证据"), (564, 350, "目前可以支持的结论")]:
    text(x, 202, w, 21, val, 13, WHITE, True)
rows = [
    ("H100 + A100 / 10.16-17", "8 项 / 4 项硬件基线", "建立独立基线；不能直接相减当配对结果"),
    ("配对回放 / 10.18", "80 × 256 固定 tokens；4 schedules", "诚实数值漂移依赖执行路径"),
    ("真实量化 / 10.18", "BF16 / NF4 / INT8；每臂 32 校准 + 64 测试", "NF4 有信号；INT8 在当前预算下较弱"),
    ("本地机制 / 10.19", "524 本地测量 + 42 HTTP + 96 SDPA 控制", "合法表示与 kernel 也改变耗时"),
    ("vLLM serving / 10.19", "600 + 600 请求；缓存 / 并发 / eager 控制", "上下文响应可复测，幅度取决于 serving"),
    ("调度追踪 / 10.19", "26 次运行；98 请求；2K / 8K token budget", "早期到达间隔包含 prefill / 调度贡献"),
]
for i, row in enumerate(rows):
    y = 225+i*32
    rect(36, y, 888, 32, WHITE if i % 2 == 0 else "#EDF2F7")
    for j, val in enumerate(row):
        text(cols[j]+10, y+8, cols[j+1]-cols[j]-18, 22, val, 11.4, INK, j == 0)
banner("下一条必须补齐的证据链：独立校准 → 合法条件误报 → 攻击检出 → 端到端成本。", size=17)


# 3. Actual quantization, exact counts and uncertainty.
new_page(
    "02 / REAL-WEIGHT QUANTIZATION",
    "真实 NF4 已有检测信号；INT8 仍是难点",
    "Qwen3-1.7B / L40S；32 条 BF16 校准；每臂 64 条测试；每条 128 tokens；名义 α = 5%。",
    ["real_quantization_l40s_20261018/pilot/summary.json", "report: real_quantization_l40s_20261018.md"],
    "图中所有比例和Wilson95%区间直接读取归档summary。阈值为32个诚实校准序列均值的最大值，"
    "严格大于阈值才拒绝。32/64等分母是序列，不是tokens。区间仅描述实现阈值下的测试计数，"
    "不包括重新校准的波动；不是三个测试联合的5%误报控制。token-DiFR需要原始共享RNG协议；"
    "token-TOPLOC是仓库内的rank-based评分，不能等同原始TOPLOC承诺协议。activation-DiFR需要内部激活。"
    "新实验不能直接声称推翻此前另一个模型/预算/协议下的NF4 0/10。",
)
text(36, 190, 615, 24, "拒绝率（BF16 为误报；NF4 / INT8 为检出）", 16, INK, True)
for x, name, color in [(36, "BF16", MUTED), (136, "NF4", TEAL), (231, "INT8", ORANGE)]:
    rect(x, 224, 12, 8, color)
    text(x+18, 219, 75, 20, name, 12, color, english=True)
plot_x, plot_w, maxval = 178, 300, 0.75
for tick in [0, .25, .5, .75]:
    xx = plot_x+tick/maxval*plot_w
    line(xx, 247, xx, 417, LINE, 0.7)
    text(xx-20, 421, 40, 18, f"{tick:.0%}", 10, MUTED, align="center", english=True)
for i, (detector, label) in enumerate([("token_difr", "token-DiFR"), ("token_toploc", "Token rank"), ("cross_entropy", "Cross-entropy")]):
    y = 252+i*58
    text(36, y+10, 137, 27, label, 14, INK, True, english=True)
    for j, (arm, color) in enumerate([("bf16", MUTED), ("nf4", TEAL), ("int8", ORANGE)]):
        rec = q["detectors"][detector]["evaluation"][arm]
        yy = y+j*15
        if rec["rate"]:
            rect(plot_x, yy, rec["rate"]/maxval*plot_w, 9, color)
        else:
            dot(plot_x, yy+4.5, 2.2, color)
        errorbar(plot_x, yy+4.5, *rec["wilson95"], plot_w, maxval, color)
        text(493, yy-3, 124, 18, f'{rec["rejected"]}/64  ({rec["rate"]:.1%})', 11, color, english=True)
rect(646, 194, 278, 240, WHITE, radius=True)
text(662, 210, 247, 24, "结果的适用边界", 18, INK, True)
text(662, 249, 247, 166,
     "DiFR 需共享 RNG 协议。\n"
     "TOPLOC 不需共享 seed，\n但 BF16 误报为 4/64。\n\n"
     "DiFR 的 0/64 误报：\n95% Wilson 上界仍为 5.66%。",
     16, MUTED, leading=1.28)
text(36, 440, 600, 15, "误差线：95% Wilson 区间。token-TOPLOC 为仓库内 rank score。", 10.2, MUTED)
banner("内部激活控制：NF4 64/64、INT8 52/64；需要 provider 额外证据，须单独报告。",
       y=458, h=28, size=15)


# 4. Optimized and eager serving plus a legitimate cache control.
new_page(
    "03 / PRODUCTION SERVING",
    "计时信号能穿过 vLLM，但依赖执行模式与流位置",
    "同一 L40S 的 loopback streaming；每轮 600 请求；延迟 18.25 分钟后用相同输入复测。",
    ["vllm_serving_l40s_20261019/summary.json", "vllm_serving_l40s_20261019.md"],
    "左图：16K请求对截为末尾2K的post-first时间比，先在六个配对文本block内取比再汇总。"
    "optimized包括compile及CUDA Graphs；eager同时关闭二者。C为客户端并发，不是内部batch size。"
    "右图：8K、C4的合法缓存cold/hit比，早期32个间隔，后期95个间隔。窗口是在pilot之后探索选择，"
    "复测沿用；不是预先独立验证的稳态窗口。误差线为六个block的min-max，不是CI。"
    "这是机制结果，没有拟合检测阈值，也没有部署FPR或TPR。",
)
text(36, 193, 535, 24, "16K full / truncated：首事件后的 ms/token 比", 16, INK, True)
text(623, 193, 301, 24, "合法 prefix cache：cold / hit", 16, INK, True)
for x, name, color in [(36, "Pilot", BLUE), (137, "Repeat", TEAL)]:
    rect(x, 226, 12, 8, color)
    text(x+18, 221, 80, 19, name, 12, color, english=True)
xp, pw, maximum = 168, 306, 3.0
for tick in [0, 1, 2, 3]:
    xx = xp+tick/maximum*pw
    line(xx, 249, xx, 411, MUTED if tick == 1 else LINE, 0.8)
    text(xx-14, 418, 28, 18, f"{tick}×", 10, MUTED, align="center", english=True)
for i, (mode, con, label) in enumerate([
    ("graph_off", 1, "Optimized C1"), ("graph_off", 4, "Optimized C4"),
    ("eager_off", 1, "Eager C1"), ("eager_off", 4, "Eager C4"),
]):
    yy = 252+i*40
    text(36, yy+4, 126, 23, label, 13, INK, english=True)
    for j, (phase, color) in enumerate([("pilot", BLUE), ("delayed_repeat", TEAL)]):
        rec = contrast(phase, mode, 16384, con)["post_first_ms_per_token"]
        y = yy+j*14
        rect(xp, y, rec["median"]/maximum*pw, 9, color)
        errorbar(xp, y+4.5, rec["min"], rec["max"], pw, maximum, color)
        text(487, y-3, 73, 18, f'{rec["median"]:.3f}×', 11, color, english=True)
rect(607, 231, 317, 195, WHITE, radius=True)
text(624, 245, 135, 20, "8K / C = 4", 13, MUTED, english=True)
text(768, 245, 67, 20, "Pilot", 12, BLUE, english=True)
text(844, 245, 66, 20, "Repeat", 12, TEAL, english=True)
for y, key, label in [(281, "early_ms_per_token", "早期 0-32"), (328, "late_ms_per_token", "后期 32-127")]:
    text(624, y, 138, 25, label, 16, INK)
    for phase, x, col in [("pilot", 768, BLUE), ("delayed_repeat", 844, TEAL)]:
        z = contrast(phase, "graph_on", 8192, 4)[key]
        text(x, y, 72, 25, f'{z["median"]:.2f}×', 18, col, True, english=True)
text(624, 377, 280, 38, "前段差异明显，后段接近 1。\n误差线为范围；全部是同机复测。", 12, MUTED)
banner("小计时差不能直接判为截断；快首 token 也可能来自合法缓存命中。", size=18)


# 5. Two honest variability sources, with their statistical units kept separate.
new_page(
    "04 / HONEST VARIABILITY",
    "合法系统变化会移动两类观测量",
    "Token：固定输出也会换分数。Clock：早期流包含调度效应。",
    ["cross_gpu_replay_20261018/summary.json", "scheduling_l40s_20261019/summary.json"],
    "左：同一H100-origin capture在A100回放，80序列×256tokens；GPU和driver同时变化。"
    "token-DiFR跨GPU每个schedule分别计算d-prime，误差线为90%paired sequence bootstrap，400 draws。"
    "这些效应不能与另一量化实验的d-prime直接相减。右：Nsight cache-cold、8K输入、C4，"
    "两个offset组的中位数；2K改8K预算后混合迭代13变4，早期间隔下降，晚期间隔近似不变。"
    "无profiler对照方向一致（16.88至13.68 ms），但每格只有一组。预算同时改变编译范围和执行shape，"
    "不是scheduler-only因果微基准；客户端时钟与trace时钟未对齐。该证据支持调度贡献，不能确定每个请求的kernel延迟。",
)
rect(36, 193, 431, 241, WHITE, radius=True)
rect(485, 193, 439, 241, WHITE, radius=True)
text(52, 208, 399, 24, "H100 → A100：诚实 token-DiFR 漂移", 17, INK, True)
text(501, 208, 403, 24, "调度预算改变：2,048 → 8,192", 17, INK, True)
x0, xw, low, high = 200, 216, -.01, .07
px = lambda z: x0+(z-low)/(high-low)*xw
line(px(0), 250, px(0), 370, LINE, 1)
for i, (key, label) in enumerate([("chunked_b1", "Batch 1 eager"), ("chunked_b4", "Batch 4"), ("sequential", "Sequential")]):
    y = 264+i*43
    rec = c["comparisons"][key]["token_difr"]
    text(52, y-8, 140, 22, label, 13, INK, english=True)
    line(px(rec["ci"][0]), y, px(rec["ci"][1]), y, BLUE, 2)
    dot(px(rec["d_prime"]), y, 4.5, BLUE)
    text(331, y+8, 108, 20, f'{rec["d_prime"]:+.4f} d′', 11, BLUE, align="right", english=True)
for tick in [0, .03, .06]:
    text(px(tick)-24, 378, 48, 19, f"{tick:.2f}", 10, MUTED, align="center", english=True)
text(52, 406, 398, 19, "90% 配对序列 bootstrap；仅一次 GPU/driver 对比", 10.7, MUTED)
records = {}
for budget in [2048, 8192]:
    runs = [r for r in sched["runs"] if r["profiler"] == "nsys" and
            r["scheduler_budget"] == budget and r["config"] == "graph_on" and
            r["arm"] == "cold" and r["concurrency"] == 4]
    assert len(runs) == 2
    records[budget] = {
        "mixed": statistics.median(sum(it["kind"] == "mixed" for it in r["scheduling"]["timeline"]) for r in runs),
        "early": statistics.median(r["arrivals"]["group_medians"]["early_ms_per_token"] for r in runs),
        "late": statistics.median(r["arrivals"]["group_medians"]["late_ms_per_token"] for r in runs),
    }
assert records[2048]["mixed"] == 13 and records[8192]["mixed"] == 4
for y, label, key, unit in [
    (256, "prefill / generation 混合迭代", "mixed", "次"),
    (302, "早期到达间隔", "early", "ms/token"),
    (348, "后期到达间隔", "late", "ms/token"),
]:
    text(501, y, 240, 24, label, 14, INK)
    a, b = records[2048][key], records[8192][key]
    value = f"{a:.0f} → {b:.0f}" if key == "mixed" else f"{a:.2f} → {b:.2f}"
    text(739, y-2, 168, 28, value, 20, TEAL, True, english=True)
    text(739, y+24, 160, 16, unit, 10, MUTED, english=key != "mixed")
text(501, 406, 403, 19, "Nsight；8K / C4 / cache cold；两组中位数", 10.7, MUTED)
banner("论文中的 honest null 必须覆盖允许的执行条件；单一参考分布不足以证明部署可靠性。", size=16)


# 6. Proposed contribution and architecture, clearly distinct from existing code.
new_page(
    "05 / PAPER THESIS",
    "已有测量框架；核心技术贡献尚未建立。",
    "最新文献已覆盖混合校准与 vLLM 集成；新主线见 BF16 技术补充 slides。",
    ["docs/VERIFICATION_SERVICE.md", "planning/NEXT_EXPERIMENTS.md", "MLSys 2027 CFP"],
    "这是论文组织与后续设计建议，不声称新联合系统已完成。现有API已有各自的token/clock检查，"
    "没有匹配校准时返回inconclusive。拟新增的serving-condition envelope、冻结的晚期窗口和joint calibration"
    "必须独立实验。合法条件集合应由可信基准实验定义，不能只信被审计provider的元数据自报。"
    "黑盒token/arrival通道、共享RNG协作通道和内部activation通道必须分别比较。相关工作应分别讨论"
    "重算检验、activation/proof方案和timing fingerprinting；仓库token_toploc只是rank scorer，"
    "不能代表原TOPLOC协议。本文的新颖性需要通过文献对照与消融证明。"
    "2026-10-20技术复审：DiFR已覆盖混合校准与vLLM集成，LLM-42和Bit-Exact进一步覆盖确定性与精确回放；"
    "本页只是实验框架，不能当作已有独立创新。详见planning/BF16_TECHNICAL_NOVELTY.md及技术补充deck。",
    status="写作建议 / 设计待验证",
)
steps = [
    ("探测请求", "长 / 短 context\n新鲜 prompt / 固定预算", TEAL_PALE, TEAL),
    ("采集证据", "tokens + SSE 到达时间\n精确记录事件合并", PALE, BLUE),
    ("适用条件 + 校准", "覆盖合法 serving 变化\n冻结统计量与阈值", ORANGE_PALE, ORANGE),
    ("输出审计结论", "异常 / 未见异常\n证据不足 → inconclusive", TEAL_PALE, TEAL),
]
for i, (title, body, fill, col) in enumerate(steps):
    x = 36+i*226
    rect(x, 196, 210, 115, fill, radius=True)
    text(x+12, 211, 186, 27, title, 18, col, True)
    text(x+12, 251, 186, 52, body, 14, INK, leading=1.4)
    if i < 3:
        text(x+212, 239, 14, 29, "›", 22, MUTED, english=True)
for x, tag, title, body, col in [
    (36, "C1 / 已有证据", "可识别性的边界", "量化、上下文与合法优化\n分别改变哪些观测量？", TEAL),
    (338, "C2 / 部分实现，待验证", "有适用条件的审计", "校准何时可用？\n何时应返回证据不足？", ORANGE),
    (640, "C3 / 待补主实验", "检出率与成本的权衡", "固定误报率下\n比较通道、预算和重算成本。", ORANGE),
]:
    text(x, 338, 284, 20, tag, 12, col, True)
    text(x, 368, 284, 25, title, 20, INK, True)
    text(x, 403, 284, 45, body, 15, MUTED, leading=1.3)
banner("技术复审：单靠上述框架不足以形成新贡献；需验证新的采样边界建模与复算收益。", y=457, h=28, size=15)


# 7. A specific writing plan, grounded in current CFP.
new_page(
    "06 / MANUSCRIPT STRUCTURE",
    "10 页正文：每一节只承担一个可验证的论点",
    "章节分配仅为框架；新的技术命题与优先实验见 BF16 技术补充 slides。",
    ["paper/paper.tex (August draft)", "mlsys.org/Conferences/2027/CallForResearchPapers"],
    "官网CFP于2026-10-20查询：10页双栏正文，不含参考文献；附录单独提交；必须有参考文献后的AI Usage Disclosure。"
    "页数分配为建议，总计10页。现稿是August 2026单H100实验汇总，不能直接作为新论文摘要。"
    "必须修改的旧叙述包括：所有攻击都可检出、普遍零成本、co-tenancy普遍单边安全以及泛化的模型身份证明。"
    "应使用最新NF4/INT8真实权重结果；历史模拟扰动和小池偏差作为方法学控制；"
    "实验室profiler仅作为解释oracle，不是黑盒验证器的输入。旧稿16x成本节省若保留需重测端到端路径，"
    "不把局部prefix scheduling速度当本设计已有收益。",
    status="写作建议",
)
head = [36, 204, 278, 612, 924]
for a, b in zip(head[:-1], head[1:]):
    rect(a, 190, b-a, 28, INK)
for x, width, value in [(46, 150, "章节"), (214, 58, "页数"), (288, 314, "必须回答的问题"), (622, 290, "主图 / 证据")]:
    text(x, 196, width, 20, value, 12, WHITE, True)
outline = [
    ("1  Introduction", "1.0", "为什么真实 serving 使验证不可靠？", "F1  合法变化与攻击的观测重叠"),
    ("2  Problem / Model", "1.0", "谁可信？可见什么？允许哪些优化？", "通道与威胁模型表；分开 RNG / 激活"),
    ("3  Characterization", "2.0", "变化来自数值、执行还是调度？", "F2  量化；F3  serving；F4  调度"),
    ("4  Design", "1.5", "怎样选统计量、校准和拒绝判定？", "架构；算法；适用条件（待验证）"),
    ("5  Evaluation", "3.5", "固定误报率下，检出率和成本如何？", "F5  TPR / FPR / cost 曲线（待补）"),
    ("6  Related / Limits", "1.0", "比已有工作多了什么？何时失败？", "相关工作对照；迁移与适应性攻击"),
]
for i, row in enumerate(outline):
    y = 218+30*i
    rect(36, y, 888, 30, WHITE if i % 2 == 0 else "#EDF2F7")
    for j, value in enumerate(row):
        text(head[j]+10, y+7, head[j+1]-head[j]-18, 21, value, 11.7, INK, j == 0)
text(36, 413, 888, 41,
     "旧稿需重写：单 H100 叙事 → 多层 serving 证据；“模型身份证明” → 有条件的行为异常检验。\n"
     "将真实量化与模拟扰动分开；将测量成本、请求成本和参考重算成本分开。", 14.2, MUTED, leading=1.4)
text(36, 465, 888, 21, "官方格式：10 页双栏；附录单独提交；参考文献后须附 AI Usage Disclosure。", 12, BLUE,
     link_url=CFP)


# 8. Prioritized research gaps, explicit falsifiers.
new_page(
    "07 / NEXT DECISIVE EXPERIMENTS",
    "先补最能改变论文结论的实验",
    "下面均为建议，尚未运行。先控制误报，再判断是否值得增加检测复杂度。",
    ["ROADMAP.md", "planning/NEXT_EXPERIMENTS.md", "MLSys 2027 research CFP"],
    "P0是第一优先级：探索数据不能既选窗口又报告独立测试精度，需冻结统计量、独立corpus/session和threshold。"
    "5%是建议operating point，并非宣称已达成；验收采用各合法条件的误报95%上界。"
    "样本数需根据区间精度和session相关性设计，不将token或并发请求当独立单位。"
    "P1同一请求同时采tokens与client time，在匹配可见性假设下比较单通道与joint方案；"
    "128/512/2048为新建议预算，不是已测量的效果。需计入探测请求、参考GPU秒、p50/p95 latency和校准成本。"
    "P2先一个额外GPU及独立corpus，再扩第二模型/较大模型。P3测试延迟填充、probe识别等适应性provider。"
    "若P0或P1失败，应收窄为实证边界/设计原则论文并诚实报告失败；不能补写部署可靠性。"
    "CFP写Oct30 20:00UTC即北京时间Oct31 04:00。主页显示PDT noon换算会差一小时，本deck明确使用CFP的UTC文本，"
    "实际提交前需再次核对OpenReview。AI协助范围包括本次归档检查、写作框架、图表脚本与slides。",
    status="待运行 / 优先级建议",
)
tasks = [
    ("P0", "独立校准 + must-not-flag", "冻结窗口；新 prompt / session；覆盖 eager、缓存、并发。",
     "判据：各合法条件 FPR 的 95% 上界 ≤ 5%。", TEAL),
    ("P1", "同请求通道消融 + 预算曲线", "token / clock / 联合；128 / 512 / 2048 tokens；报告重算与请求成本。",
     "判据：相同 FPR 下，TPR 或成本确有改善；否则删减设计。", BLUE),
    ("P2", "跨条件迁移 + 实际量化扩展", "新语料、额外 GPU / 模型；匹配真实 NF4 / INT8，随后 AWQ / GPTQ。",
     "判据：区分阈值迁移成功与需重新校准的条件。", BLUE),
    ("P3", "适应性 provider", "加入延迟填充与 probe 识别；允许输出 inconclusive。",
     "判据：明确可绕过的范围与审计适用边界。", ORANGE),
]
for i, (priority, title, design, gate, col) in enumerate(tasks):
    y = 190+i*65
    rect(36, y, 888, 60, WHITE, radius=True)
    text(49, y+17, 46, 25, priority, 19, col, True, english=True)
    text(106, y+7, 276, 25, title, 16, INK, True)
    text(394, y+8, 514, 22, design, 11.8, MUTED)
    text(394, y+33, 514, 22, gate, 11.8, col)
text(36, 461, 888, 22,
     "MLSys 2027 CFP：10 月 30 日 20:00 UTC 截稿（北京时间 10 月 31 日 04:00）。", 13, BLUE,
     link_url=CFP)


pdf.save()
prs.save(PPTX)
evidence["derived_scheduler_cold_medians"] = records
evidence["conference"] = {
    "url": CFP, "checked": "2026-10-20", "body_page_limit": 10,
    "deadline_utc_per_cfp": "2026-10-30T20:00:00Z",
    "note": "Conference homepage displays noon PDT, one hour earlier than CFP's UTC deadline; recheck submission portal.",
}
(QA / "evidence.json").write_text(json.dumps(evidence, indent=2, ensure_ascii=False)+"\n")
(QA / "slide-index.json").write_text(json.dumps(pages, indent=2, ensure_ascii=False)+"\n")
(QA / "layout-validation.json").write_text(json.dumps({
    "slides": len(pages), "overflow_issues": issues,
    "renderer": "shared coordinates; editable PPTX objects and ReportLab vector PDF",
    "pptx_rendering_limit": "PowerPoint/Keynote rendering not checked; PDF inspected separately.",
}, indent=2, ensure_ascii=False)+"\n")
if issues:
    raise SystemExit("\n".join(issues))
print(f"Built {len(pages)} slides with no measured text overflow.\n{PPTX}\n{PDF}")
